"""Ekspor briefing pimpinan dan masukan kebijakan. Hanya saran yang punya tautan bukti."""

from __future__ import annotations

import html
import io

from pptx import Presentation
from pptx.util import Inches, Pt
from sqlalchemy.orm import Session

from app.services.dss import DISCLAIMER, list_advices


def briefing_html(db: Session) -> str:
    payload = list_advices(db, {})
    items = [item for item in payload["items"] if item["status"] not in {"kedaluwarsa", "ditolak"} and item["evidence"]]
    blocks = []
    for item in items:
        proof = "".join(
            f"<li>{html.escape(ev['source_name'])}: {html.escape(ev['title'])} — "
            f"<a href=\"{html.escape(ev['url'], quote=True)}\">Buka sumber asli</a></li>"
            for ev in item["evidence"]
        )
        blocks.append(
            f"<section><h2>{html.escape(item['kind_label'])}: {html.escape(item['title'])}</h2>"
            f"<p>Urgensi {html.escape(item['urgency'])} · status {html.escape(item['status'])} · keyakinan {html.escape(item['confidence'])}</p>"
            f"<p>{html.escape(item['situation'])}</p>"
            f"<p><strong>Rekomendasi.</strong> {html.escape(item['recommendation'])}</p>"
            f"<ul>{proof}</ul></section>"
        )
    body = "\n".join(blocks) or "<p>Belum ada saran dengan bukti.</p>"
    return f"""<!doctype html><html lang="id"><head><meta charset="utf-8"><title>Briefing Pimpinan</title>
<style>body{{font-family:sans-serif;margin:2rem;color:#0B1F3A}} a{{color:#1E4A7A}}</style></head>
<body><h1>Briefing Pimpinan</h1><p>Diperbarui: {html.escape(payload.get('refreshed_at') or '')} WIB</p>
<p>{html.escape(DISCLAIMER)}</p>{body}</body></html>"""


def policy_html(db: Session, topic: str = "") -> str:
    payload = list_advices(db, {"kind": "kebijakan"})
    items = payload["items"]
    if topic:
        items = [item for item in items if topic.casefold() in f"{item['policy_topic']} {item['title']}".casefold()]
    blocks = []
    for item in items:
        if not item["evidence"]:
            continue
        proof = "".join(
            f"<li>{html.escape(ev['source_name'])}: {html.escape(ev['title'])} — "
            f"<a href=\"{html.escape(ev['url'], quote=True)}\">Buka sumber asli</a></li>"
            for ev in item["evidence"]
        )
        blocks.append(f"<section><h2>{html.escape(item['title'])}</h2><p>{html.escape(item['situation'])}</p><ul>{proof}</ul></section>")
    body = "\n".join(blocks) or "<p>Data belum cukup untuk topik ini.</p>"
    return f"""<!doctype html><html lang="id"><head><meta charset="utf-8"><title>Masukan Kebijakan</title></head>
<body><h1>Masukan kebijakan — persepsi publik</h1><p>{html.escape(DISCLAIMER)}</p>
<p>Bukan rekomendasi kebijakan moneter.</p>{body}</body></html>"""


def briefing_pptx(db: Session) -> bytes:
    payload = list_advices(db, {})
    items = [item for item in payload["items"] if item["status"] not in {"kedaluwarsa", "ditolak"} and item["evidence"]][:8]
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    def add(title: str, lines: list[str]) -> None:
        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(6.6))
        frame = box.text_frame
        frame.word_wrap = True
        head = frame.paragraphs[0]
        head.text = title
        head.font.size = Pt(24)
        head.font.bold = True
        for line in lines:
            paragraph = frame.add_paragraph()
            paragraph.text = line[:500]
            paragraph.font.size = Pt(14)

    add("Briefing Pimpinan", [f"Saran diperbarui: {payload.get('refreshed_at') or '-'} WIB", DISCLAIMER, "Tidak ada tindakan yang dikirim otomatis."])
    for item in items:
        links = [f"{ev['source_name']}: {ev['title']} — {ev['url']}" for ev in item["evidence"][:4]]
        add(
            item["title"][:80],
            [f"{item['kind_label']} · urgensi {item['urgency']} · {item['status']}", item["situation"][:400], item["recommendation"][:400], *links],
        )
    if not items:
        add("Belum ada saran", ["Tidak ada saran dengan bukti sumber pada siklus ini."])
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
