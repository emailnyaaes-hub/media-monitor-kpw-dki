"""Ekspor laporan. PDF dilakukan lewat cetak peramban agar tidak menambah mesin berat."""

from __future__ import annotations

import csv
import html
import io
from datetime import datetime

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from sqlalchemy.orm import Session

from app.constants import DISCLAIMER
from app.services import analytics


def narrative(db: Session, window: dict) -> dict:
    snap = analytics.overview(db, window)
    issues = analytics.list_issues(db, window)[:5]
    kpis = snap["kpis"]
    net = kpis["net_sentiment"]["value"]
    tone = "cenderung terjaga" if net >= 10 else "campuran" if net >= -10 else "perlu perhatian"
    lines = [
        f"Pada {snap['period']['from']} sampai {snap['period']['to']} tercatat {kpis['total_mention']['value']} mention.",
        f"Sentimen bersih {net} (persen positif dikurangi persen negatif), sehingga persepsi {tone}.",
        f"Isu kritis yang masih terbuka: {kpis['critical_issues']['value']}.",
    ]
    if issues:
        lines.append("Isu dengan skor risiko tertinggi: " + "; ".join(item["title"] for item in issues[:3]) + ".")
    lines.append("Jangkauan platform tidak tersedia dari RSS. Angka 0 berarti metrik itu tidak dikumpulkan, bukan nol pembaca.")
    lines.append("Ringkasan ini hanya memakai judul dan cuplikan yang berhasil dikumpulkan. Buka tautan sumber sebelum mengutip.")
    sources = [
        {
            "id": item["id"],
            "title": item["title"],
            "source_name": item["source_name"],
            "url": item.get("url", ""),
            "published_at": item["published_at"],
            "quotes_bi": bool(item.get("quotes_bi")),
        }
        for item in analytics.search_mentions(db, window, {"page": 1, "page_size": 8, "sort": "terbaru"})["items"]
    ]
    return {
        "title": "Ringkasan eksekutif",
        "generated_at": datetime.now().isoformat(timespec="minutes"),
        "paragraphs": lines,
        "disclaimer": DISCLAIMER,
        "overview": snap,
        "issues": issues,
        "sources": sources,
        "policy_feedback": analytics.policy_feedback(db, window),
    }


def workbook(db: Session, window: dict) -> bytes:
    report = narrative(db, window)
    book = Workbook()
    sheet = book.active
    sheet.title = "Ringkasan"
    sheet.append(["Media Monitor KPw BI DKI Jakarta"])
    sheet.append(["Mode", "Sumber asli"])
    sheet.append(["Periode", report["overview"]["period"]["from"], report["overview"]["period"]["to"]])
    for paragraph in report["paragraphs"]:
        sheet.append([paragraph])
    sheet.append([])
    sheet.append([report["disclaimer"]])

    issues = book.create_sheet("Isu")
    issues.append(["Kode", "Judul", "Kategori", "Stance", "Risiko", "Relevansi kebijakan", "Kecepatan", "Dampak", "Kuadran", "Status", "Jumlah mention"])
    for item in report["issues"]:
        issues.append([
            item["code"], item["title"], item["category"], item["stance"], item["risk_score"],
            item["policy_relevance"], item["spread_velocity"], item["impact_level"], item["quadrant"],
            item["status"], item["mention_count"],
        ])

    policy = book.create_sheet("Umpan kebijakan")
    policy.append(["Kebijakan", "Mention", "Positif %", "Netral %", "Negatif %", "Aspirasi", "Kritik"])
    for item in report["policy_feedback"]:
        policy.append([
            item["policy"], item["mentions"], item["positif"], item["netral"], item["negatif"],
            " | ".join(f"{quote['title']} {quote['url']}" for quote in item["aspirations"]),
            " | ".join(f"{quote['title']} {quote['url']}" for quote in item["criticisms"]),
        ])

    mentions = book.create_sheet("Mention")
    mentions.append(["Waktu", "Sumber", "Platform", "Judul", "Sentimen", "Stance", "Risiko", "Kategori", "Status", "Mengutip BI", "URL"])
    rows = analytics.search_mentions(db, window, {"page": 1, "page_size": 1000, "sort": "terbaru"})
    for item in rows["items"]:
        mentions.append([
            item["published_at"], item["source_name"], item["platform"], item["title"],
            item["sentiment"], item["stance"], item["risk_score"], item["category"], item["status"],
            "ya" if item.get("quotes_bi") else "", item.get("url", ""),
        ])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def csv_mentions(db: Session, window: dict, params: dict) -> str:
    params = {**params, "page": 1, "page_size": 1000}
    rows = analytics.search_mentions(db, window, params)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["id", "waktu", "sumber", "platform", "judul", "sentimen", "stance", "risiko", "kategori", "status", "jangkauan", "mengutip_bi", "url"])
    for item in rows["items"]:
        writer.writerow([
            item["id"], item["published_at"], item["source_name"], item["platform"], item["title"],
            item["sentiment"], item["stance"], item["risk_score"], item["category"], item["status"],
            item["reach_estimate"], "ya" if item.get("quotes_bi") else "", item.get("url", ""),
        ])
    return buffer.getvalue()


def pptx_report(db: Session, window: dict) -> bytes:
    """Beberapa salindia ringkas. Ini bahan rapat, bukan keputusan."""
    report = narrative(db, window)
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    def add_slide(title: str, lines: list[str]) -> None:
        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(6.5))
        frame = box.text_frame
        frame.word_wrap = True
        head = frame.paragraphs[0]
        head.text = title
        head.font.size = Pt(28)
        head.font.bold = True
        for line in lines:
            paragraph = frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(16)
            paragraph.level = 0

    kpis = report["overview"]["kpis"]
    period = report["overview"]["period"]
    add_slide(
        "Laporan Media Monitor",
        [
            f"KPw BI DKI Jakarta · {period['from']} s.d. {period['to']}",
            "Hanya konten pihak luar. Kanal resmi Bank Indonesia tidak dihitung. Angka wajib dibaca dari tautan sumber.",
            *report["paragraphs"],
        ],
    )
    add_slide(
        "Angka utama",
        [
            f"Mention: {kpis['total_mention']['value']}",
            f"Sentimen bersih: {kpis['net_sentiment']['value']}",
            f"Isu kritis terbuka: {kpis['critical_issues']['value']}",
            f"Jangkauan estimasi: {kpis['reach']['value']}",
        ],
    )
    add_slide(
        "Isu yang perlu perhatian",
        [
            f"{item['title']} — {item['stance']}, risiko {item['risk_score']}, {item['quadrant']}"
            for item in report["issues"]
        ] or ["Tidak ada isu pada rentang ini."],
    )
    sources = analytics.search_mentions(db, window, {"page": 1, "page_size": 8, "sort": "terbaru"})["items"]
    add_slide(
        "Sumber yang dikutip",
        [f"{item['source_name']}: {item['title']} — {item.get('url', '')}" for item in sources] or ["Belum ada sumber pada rentang ini."],
    )
    add_slide("Catatan", [report["disclaimer"]])
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def html_report(db: Session, window: dict) -> str:
    report = narrative(db, window)
    kpis = report["overview"]["kpis"]
    issue_rows = "".join(
        "<tr>"
        f"<td>{item['title']}</td><td>{item['stance']}</td><td>{item['risk_score']}</td>"
        f"<td>{item['quadrant']}</td><td>{item['status']}</td>"
        "</tr>"
        for item in report["issues"]
    )
    source_rows = analytics.search_mentions(db, window, {"page": 1, "page_size": 12, "sort": "terbaru"})["items"]
    source_html = "".join(
        f"<li>{html.escape(item['source_name'])}: {html.escape(item['title'])} — <a href=\"{html.escape(item.get('url', ''), quote=True)}\">Buka sumber asli</a></li>"
        for item in source_rows
    )
    paragraphs = "".join(f"<p>{line}</p>" for line in report["paragraphs"])
    return f"""<!DOCTYPE html>
<html lang="id">
<head>
  <meta charset="utf-8" />
  <title>Laporan Media Monitor</title>
  <style>
    body {{ font-family: "Plus Jakarta Sans", "Segoe UI", sans-serif; color: #0B1F3A; margin: 32px; }}
    h1 {{ font-size: 22px; margin-bottom: 4px; }}
    .gold {{ height: 4px; background: #C4A35A; width: 96px; margin: 8px 0 16px; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
    th, td {{ border-bottom: 1px solid #E4E8EE; text-align: left; padding: 8px; vertical-align: top; }}
    .note {{ background: #F8F1DE; padding: 12px; font-size: 12px; }}
    .kpi {{ display: flex; gap: 16px; margin: 16px 0; }}
    .kpi div {{ border: 1px solid #E4E8EE; padding: 10px 14px; min-width: 120px; }}
    @media print {{ .noprint {{ display: none; }} }}
  </style>
</head>
<body>
  <button class="noprint" onclick="window.print()">Cetak atau simpan PDF</button>
  <h1>Laporan Media Monitor</h1>
  <div>KPw BI Provinsi DKI Jakarta · Sumber asli, cuplikan singkat</div>
  <div class="gold"></div>
  <div class="kpi">
    <div><strong>{kpis['total_mention']['value']}</strong><br>Mention</div>
    <div><strong>{kpis['net_sentiment']['value']}</strong><br>Sentimen bersih</div>
    <div><strong>{kpis['critical_issues']['value']}</strong><br>Isu kritis</div>
  </div>
  {paragraphs}
  <h2>Isu utama</h2>
  <table>
    <thead><tr><th>Isu</th><th>Stance</th><th>Risiko</th><th>Kuadran</th><th>Status</th></tr></thead>
    <tbody>{issue_rows}</tbody>
  </table>
  <h2>Sumber</h2>
  <ul>{source_html or "<li>Belum ada sumber pada rentang ini.</li>"}</ul>
  <p class="note">{report['disclaimer']}</p>
</body>
</html>"""
